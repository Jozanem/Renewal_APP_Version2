from docx.shared import Pt, RGBColor, Inches


def set_font_style(
    paragraph,
    bold=False,
    underline=False,
    color=RGBColor(0, 0, 0),
    size=11,
):
    """
    Apply Cambria styling to a paragraph.
    """

    for run in paragraph.runs:
        run.font.name = "Cambria"
        run.font.size = Pt(size)
        run.bold = bold
        run.underline = underline
        run.font.color.rgb = color


def add_header_footer_first_page(doc, header_image_path, footer_image_path):
    """
    Add header + footer ONLY on first page.
    """

    section = doc.sections[0]

    # ⭐ IMPORTANT
    section.different_first_page_header_footer = True

    # HEADER FIRST PAGE
    header = section.first_page_header
    header_para = header.add_paragraph()
    header_run = header_para.add_run()
    header_run.add_picture(header_image_path, width=Inches(6))
    header_para.alignment = 1

    # FOOTER FIRST PAGE
    footer = section.first_page_footer
    footer_para = footer.add_paragraph()
    footer_run = footer_para.add_run()
    footer_run.add_picture(footer_image_path, width=Inches(6))
    footer_para.alignment = 1
